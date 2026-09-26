from __future__ import annotations

import math
import mimetypes
import re
import shutil
import subprocess
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pypdf import PdfReader


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".log"}
CAD_EXTENSIONS = {".dxf", ".dwg"}
CAD_MIME_TYPES = {
    ".dxf": "image/vnd.dxf",
    ".dwg": "image/vnd.dwg",
}
MAX_DWG_STRING_SCAN_BYTES = 8 * 1024 * 1024


@dataclass(frozen=True)
class TextChunk:
    text: str
    page: int | None
    section: str


@dataclass(frozen=True)
class VisualItem:
    image_path: Path
    page: int | None
    label: str


def detect_content_type(path: Path, fallback: str = "application/octet-stream") -> str:
    if path.suffix.lower() in CAD_MIME_TYPES:
        return CAD_MIME_TYPES[path.suffix.lower()]
    return mimetypes.guess_type(path.name)[0] or fallback


def extract_text_chunks(path: Path, work_dir: Path | None = None) -> tuple[list[TextChunk], int]:
    extension = path.suffix.lower()
    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".docx":
        return _extract_docx(path), 0
    if extension in TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return chunk_text(text, None, "documento"), 0
    if extension == ".dxf":
        return _extract_dxf(path), 1
    if extension == ".dwg":
        return _extract_dwg(path, work_dir), 1
    return [], 0


def _extract_pdf(path: Path) -> tuple[list[TextChunk], int]:
    reader = PdfReader(str(path))
    chunks: list[TextChunk] = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        chunks.extend(chunk_text(text, page_index, f"page {page_index}"))
    return chunks, len(reader.pages)


def _extract_docx(path: Path) -> list[TextChunk]:
    from docx import Document

    document = Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return chunk_text("\n\n".join(paragraphs), None, "documento")


def _extract_dxf(path: Path, source_label: str | None = None) -> list[TextChunk]:
    try:
        import ezdxf
    except ImportError:
        return chunk_text(
            f"CAD drawing: {path.name}\nDXF parsing requires the ezdxf Python package.",
            None,
            "cad metadata",
        )

    document = ezdxf.readfile(path)
    summary = _summarize_dxf(document, path.name, source_label)
    return chunk_text(summary, None, "cad drawing")


def _extract_dwg(path: Path, work_dir: Path | None) -> list[TextChunk]:
    converted = convert_dwg_to_dxf(path, work_dir)
    if converted:
        chunks = _extract_dxf(converted, f"Converted from DWG: {path.name}")
        if chunks:
            return chunks
    return _extract_dwg_fallback(path)


def _summarize_dxf(document, filename: str, source_label: str | None = None) -> str:
    layers = _unique_limited(layer.dxf.name for layer in document.layers if layer.dxf.name)
    layouts = _unique_limited(layout.name for layout in document.layouts if layout.name)
    blocks = _unique_limited(block.name for block in document.blocks if not block.name.startswith("*"))
    entity_counts: Counter[str] = Counter()
    text_items: list[str] = []
    block_refs: list[str] = []

    for entity in _iter_dxf_entities(document):
        try:
            entity_type = entity.dxftype()
            entity_counts[entity_type] += 1
            text = _entity_text(entity)
            if text:
                text_items.append(text)
            if entity_type == "INSERT":
                name = getattr(entity.dxf, "name", "")
                layer = getattr(entity.dxf, "layer", "")
                if name:
                    block_refs.append(f"{name} on layer {layer}".strip())
                for attrib in getattr(entity, "attribs", []):
                    attrib_text = _clean_cad_text(getattr(attrib.dxf, "text", ""))
                    if attrib_text:
                        tag = getattr(attrib.dxf, "tag", "")
                        text_items.append(f"{tag}: {attrib_text}" if tag else attrib_text)
        except Exception:
            continue

    lines = [
        f"CAD drawing: {filename}",
        f"Source: {source_label}" if source_label else "",
        f"DXF version: {document.header.get('$ACADVER', 'unknown')}",
        f"Drawing units code: {document.header.get('$INSUNITS', 'unknown')}",
        f"Layers: {', '.join(layers) if layers else 'none detected'}",
        f"Layouts: {', '.join(layouts) if layouts else 'none detected'}",
        f"Blocks: {', '.join(blocks) if blocks else 'none detected'}",
        f"Entity counts: {_format_counter(entity_counts)}",
    ]

    unique_block_refs = _unique_limited(block_refs, limit=80)
    if unique_block_refs:
        lines.append("Block references:")
        lines.extend(unique_block_refs)

    unique_text_items = _unique_limited(text_items, limit=300)
    if unique_text_items:
        lines.append("Text and annotations:")
        lines.extend(unique_text_items)

    return "\n".join(line for line in lines if line)


def _extract_dwg_fallback(path: Path) -> list[TextChunk]:
    strings = _extract_binary_strings(path, limit=220)
    header = _dwg_header(path)
    size_mb = path.stat().st_size / (1024 * 1024)
    lines = [
        f"CAD drawing: {path.name}",
        "Format: DWG",
        f"DWG header: {header or 'unknown'}",
        f"File size: {size_mb:.2f} MB",
        "DWG is a proprietary binary CAD format. Install ODA File Converter or LibreDWG to enable full local DWG to DXF conversion.",
    ]
    if strings:
        lines.append("Readable embedded strings found in the DWG:")
        lines.extend(strings)
    return chunk_text("\n".join(lines), None, "cad dwg metadata")


def chunk_text(text: str, page: int | None, section: str, max_chars: int = 1200) -> list[TextChunk]:
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not cleaned:
        return []

    paragraphs = cleaned.split("\n")
    chunks: list[TextChunk] = []
    current: list[str] = []
    current_len = 0

    for paragraph in paragraphs:
        paragraph_len = len(paragraph)
        if current and current_len + paragraph_len + 1 > max_chars:
            chunks.append(TextChunk(text="\n".join(current), page=page, section=section))
            current = []
            current_len = 0
        current.append(paragraph)
        current_len += paragraph_len + 1

    if current:
        chunks.append(TextChunk(text="\n".join(current), page=page, section=section))

    return chunks


def build_visual_items(
    path: Path,
    preview_dir: Path,
    max_pdf_pages: int,
    display_name: str | None = None,
) -> tuple[list[VisualItem], Path | None]:
    extension = path.suffix.lower()
    if extension in IMAGE_EXTENSIONS:
        preview = _make_image_preview(path, preview_dir / f"{path.stem}.jpg")
        label = display_name or path.name
        return [VisualItem(preview, None, label)] if preview else [], preview
    if extension == ".pdf":
        return _render_pdf_pages(path, preview_dir, max_pdf_pages)
    if extension in CAD_EXTENSIONS:
        return _render_cad_preview(path, preview_dir, display_name)
    return [], None


def _make_image_preview(source: Path, destination: Path, max_size: tuple[int, int] = (1300, 1300)) -> Path | None:
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(source) as image:
            image.thumbnail(max_size)
            if image.mode not in {"RGB", "L"}:
                image = image.convert("RGB")
            image.save(destination, "JPEG", quality=86, optimize=True)
        return destination
    except Exception:
        return None


def _render_pdf_pages(path: Path, preview_dir: Path, max_pages: int) -> tuple[list[VisualItem], Path | None]:
    try:
        import pypdfium2 as pdfium
    except Exception:
        return [], None

    visual_items: list[VisualItem] = []
    first_preview: Path | None = None
    document = pdfium.PdfDocument(str(path))
    page_count = min(len(document), max_pages)

    for page_index in range(page_count):
        page = document[page_index]
        bitmap = page.render(scale=1.25).to_pil()
        preview_path = preview_dir / f"{path.stem}-page-{page_index + 1}.jpg"
        preview_path.parent.mkdir(parents=True, exist_ok=True)
        if bitmap.mode not in {"RGB", "L"}:
            bitmap = bitmap.convert("RGB")
        bitmap.thumbnail((1300, 1300))
        bitmap.save(preview_path, "JPEG", quality=84, optimize=True)

        first_preview = first_preview or preview_path
        visual_items.append(VisualItem(preview_path, page_index + 1, f"page {page_index + 1}"))

    return visual_items, first_preview


def _render_cad_preview(path: Path, preview_dir: Path, display_name: str | None = None) -> tuple[list[VisualItem], Path | None]:
    label = display_name or path.name
    source_path = path
    if path.suffix.lower() == ".dwg":
        source_path = convert_dwg_to_dxf(path, preview_dir) or path

    preview_path = preview_dir / f"{path.stem}.jpg"
    if source_path.suffix.lower() == ".dxf":
        preview = _render_dxf_preview(source_path, preview_path, label)
    else:
        preview = _render_related_pdf_cad_preview(path, preview_path, label) or _make_cad_placeholder_preview(
            path,
            preview_path,
            label,
        )
    return ([VisualItem(preview, None, label)] if preview else []), preview


def _render_dxf_preview(path: Path, destination: Path, label: str) -> Path | None:
    try:
        import ezdxf
    except ImportError:
        return _make_cad_placeholder_preview(path, destination, label, ["DXF preview requires ezdxf."])

    try:
        document = ezdxf.readfile(path)
        shapes, labels = _collect_preview_items(document)
    except Exception:
        return _make_cad_placeholder_preview(path, destination, label, ["DXF preview could not be rendered."])

    if not shapes and not labels:
        return _make_cad_placeholder_preview(path, destination, label, ["No drawable CAD entities were found."])

    destination.parent.mkdir(parents=True, exist_ok=True)
    image_size = 1300
    margin = 78
    image = Image.new("RGB", (image_size, image_size), "#f8fafc")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default()

    map_point = _make_iso_mapper(shapes, labels, image_size, margin)
    _draw_cad_stage(draw, image_size)
    draw.text((48, 44), label[:120], fill="#0f172a", font=font)

    depth = 20
    for shape in shapes[:12000]:
        points = [map_point(point) for point in shape.points]
        if len(points) < 2:
            continue
        if shape.closed and len(points) > 2:
            points.append(points[0])
        back_points = [(point[0] + depth, point[1] + depth) for point in points]
        draw.line(back_points, fill="#94a3b8", width=3, joint="curve")

    for shape in shapes[:12000]:
        points = [map_point(point) for point in shape.points]
        if len(points) < 2:
            continue
        color = _layer_color(shape.layer)
        if shape.closed and len(points) > 2:
            points.append(points[0])
            _draw_shape_sides(draw, points, depth, "#cbd5e1")
        draw.line(points, fill=color, width=2, joint="curve")

    for item in labels[:160]:
        point = map_point(item.point)
        text = item.text[:80]
        draw.ellipse((point[0] + depth - 3, point[1] + depth - 3, point[0] + depth + 3, point[1] + depth + 3), fill="#94a3b8")
        draw.ellipse((point[0] - 3, point[1] - 3, point[0] + 3, point[1] + 3), fill="#0f766e")
        draw.text((point[0] + 8, point[1] - 7), text, fill="#334155", font=font)

    image.save(destination, "JPEG", quality=88, optimize=True)
    return destination


def _make_iso_mapper(shapes: list["CadShape"], labels: list["CadLabel"], image_size: int, margin: int):
    min_x, min_y, max_x, max_y = _preview_bounds(shapes, labels)
    raw_points = [
        _iso_raw(point, min_x, min_y, max_y)
        for shape in shapes
        for point in shape.points
    ]
    raw_points.extend(_iso_raw(label.point, min_x, min_y, max_y) for label in labels)

    raw_min_x = min(point[0] for point in raw_points)
    raw_min_y = min(point[1] for point in raw_points)
    raw_max_x = max(point[0] for point in raw_points)
    raw_max_y = max(point[1] for point in raw_points)
    span_x = max(raw_max_x - raw_min_x, 1.0)
    span_y = max(raw_max_y - raw_min_y, 1.0)
    scale = min((image_size - margin * 2) / span_x, (image_size - margin * 2) / span_y)
    drawn_width = span_x * scale
    drawn_height = span_y * scale
    offset_x = (image_size - margin * 2 - drawn_width) / 2
    offset_y = (image_size - margin * 2 - drawn_height) / 2 + 18

    def map_point(point: tuple[float, float]) -> tuple[float, float]:
        raw_x, raw_y = _iso_raw(point, min_x, min_y, max_y)
        return (
            margin + offset_x + (raw_x - raw_min_x) * scale,
            margin + offset_y + (raw_y - raw_min_y) * scale,
        )

    return map_point


def _iso_raw(point: tuple[float, float], min_x: float, min_y: float, max_y: float) -> tuple[float, float]:
    x, y = point
    local_x = x - min_x
    local_y = y - min_y
    inverted_y = max_y - y
    return local_x + local_y * 0.42, inverted_y * 0.58 + local_x * 0.16


def _draw_cad_stage(draw: ImageDraw.ImageDraw, image_size: int) -> None:
    draw.rectangle((32, 32, image_size - 32, image_size - 32), outline="#cbd5e1", width=2)
    floor = [(96, 1034), (1128, 1034), (1214, 1120), (180, 1120)]
    draw.polygon(floor, fill="#e2e8f0", outline="#cbd5e1")
    for index in range(12):
        x = 116 + index * 84
        draw.line((x, 1040, x + 88, 1120), fill="#cbd5e1", width=1)
    for index in range(5):
        y = 1052 + index * 16
        draw.line((108, y, 1200, y), fill="#cbd5e1", width=1)


def _draw_shape_sides(draw: ImageDraw.ImageDraw, points: list[tuple[float, float]], depth: int, color: str) -> None:
    for start, end in zip(points, points[1:], strict=False):
        side = [start, end, (end[0] + depth, end[1] + depth), (start[0] + depth, start[1] + depth)]
        draw.polygon(side, fill=color)


def _render_related_pdf_cad_preview(path: Path, destination: Path, label: str) -> Path | None:
    pdf_path = _find_related_pdf(path)
    if not pdf_path:
        return None
    try:
        page_image = _render_pdf_first_page_image(pdf_path)
    except Exception:
        return None
    return _make_pdf_card_preview(page_image, destination, label, f"DWG preview from paired PDF: {pdf_path.name}")


def _find_related_pdf(path: Path) -> Path | None:
    tokens = _cad_match_tokens(path.name)
    pdfs = sorted(path.parent.glob("*.pdf")) + sorted(path.parent.glob("*.PDF"))
    if not pdfs:
        return None
    for token in tokens:
        for pdf in pdfs:
            if token.lower() in pdf.name.lower():
                return pdf
    return pdfs[0] if len(pdfs) == 1 else None


def _cad_match_tokens(name: str) -> list[str]:
    tokens = []
    for match in re.finditer(r"(?<!\d)(\d{5,6}(?:_\d+)?)(?!\d)", name):
        token = match.group(1)
        tokens.append(token)
        tokens.append(token.split("_", 1)[0])
    return _unique_limited(tokens, limit=6)


def _render_pdf_first_page_image(path: Path) -> Image.Image:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    page = document[0]
    try:
        image = page.render(scale=1.5).to_pil()
        if image.mode != "RGB":
            image = image.convert("RGB")
        image.thumbnail((900, 980))
        return image
    finally:
        close_page = getattr(page, "close", None)
        close_document = getattr(document, "close", None)
        if close_page:
            close_page()
        if close_document:
            close_document()


def _make_pdf_card_preview(page_image: Image.Image, destination: Path, label: str, note: str) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (1300, 900), "#f8fafc")
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.load_default()
    _draw_preview_grid(draw)

    page = page_image.copy()
    page.thumbnail((820, 680))
    bordered = Image.new("RGB", (page.width + 28, page.height + 28), "#e2e8f0")
    bordered.paste(page, (14, 14))
    rotated = bordered.rotate(-8, resample=Image.Resampling.BICUBIC, expand=True, fillcolor="#f8fafc")
    mask = Image.new("L", rotated.size, 0)
    mask_draw = ImageDraw.Draw(mask)
    mask_draw.rectangle((12, 12, rotated.width - 12, rotated.height - 12), fill=255)
    mask = mask.rotate(-8, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=0).resize(rotated.size)

    x = (canvas.width - rotated.width) // 2
    y = 138
    shadow = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    shadow_mask = Image.new("L", rotated.size, 150)
    shadow_mask = shadow_mask.filter(ImageFilter.GaussianBlur(18))
    shadow.paste((15, 23, 42, 95), (x + 34, y + 44), shadow_mask)
    canvas = Image.alpha_composite(canvas.convert("RGBA"), shadow).convert("RGB")
    draw = ImageDraw.Draw(canvas)

    thickness = [(x + 44, y + rotated.height - 8), (x + rotated.width - 22, y + rotated.height - 8), (x + rotated.width - 54, y + rotated.height + 34), (x + 12, y + rotated.height + 34)]
    draw.polygon(thickness, fill="#cbd5e1", outline="#94a3b8")
    canvas.paste(rotated, (x, y))

    draw.rectangle((42, 40, 1258, 112), fill="#ffffff", outline="#cbd5e1")
    draw.text((64, 60), label[:135], fill="#0f172a", font=font)
    draw.text((64, 86), note[:145], fill="#475569", font=font)
    draw.text((64, 830), "Simulated 3D technical preview", fill="#0f766e", font=font)
    canvas.save(destination, "JPEG", quality=88, optimize=True)
    return destination


def _draw_preview_grid(draw: ImageDraw.ImageDraw) -> None:
    floor = [(96, 790), (1110, 790), (1232, 884), (210, 884)]
    draw.polygon(floor, fill="#e2e8f0", outline="#cbd5e1")
    for index in range(13):
        x = 116 + index * 76
        draw.line((x, 794, x + 118, 884), fill="#cbd5e1", width=1)
    for index in range(6):
        y = 804 + index * 15
        draw.line((106, y, 1220, y), fill="#cbd5e1", width=1)


@dataclass
class CadShape:
    points: list[tuple[float, float]]
    layer: str
    closed: bool = False


@dataclass
class CadLabel:
    text: str
    point: tuple[float, float]
    layer: str


def _collect_preview_items(document) -> tuple[list[CadShape], list[CadLabel]]:
    shapes: list[CadShape] = []
    labels: list[CadLabel] = []
    for entity in _iter_dxf_entities(document, limit=25000):
        _append_preview_entity(entity, shapes, labels)
    return shapes, labels


def _append_preview_entity(entity, shapes: list[CadShape], labels: list[CadLabel], depth: int = 0) -> None:
    try:
        entity_type = entity.dxftype()
        layer = getattr(entity.dxf, "layer", "0") or "0"
        if entity_type == "INSERT" and depth < 2:
            for virtual_entity in entity.virtual_entities():
                _append_preview_entity(virtual_entity, shapes, labels, depth + 1)
            return
        if entity_type == "LINE":
            shapes.append(CadShape([_xy(entity.dxf.start), _xy(entity.dxf.end)], layer))
        elif entity_type == "LWPOLYLINE":
            points = [(float(point[0]), float(point[1])) for point in entity.get_points("xy")]
            if len(points) > 1:
                shapes.append(CadShape(points, layer, bool(entity.closed)))
        elif entity_type == "POLYLINE":
            points = [_xy(vertex.dxf.location) for vertex in entity.vertices]
            if len(points) > 1:
                shapes.append(CadShape(points, layer, bool(entity.is_closed)))
        elif entity_type == "CIRCLE":
            shapes.append(CadShape(_circle_points(_xy(entity.dxf.center), float(entity.dxf.radius)), layer, True))
        elif entity_type == "ARC":
            shapes.append(
                CadShape(
                    _arc_points(
                        _xy(entity.dxf.center),
                        float(entity.dxf.radius),
                        float(entity.dxf.start_angle),
                        float(entity.dxf.end_angle),
                    ),
                    layer,
                )
            )
        elif entity_type in {"TEXT", "MTEXT", "ATTRIB", "ATTDEF", "DIMENSION"}:
            text = _entity_text(entity)
            point = _entity_insert_point(entity)
            if text and point:
                labels.append(CadLabel(text, point, layer))
    except Exception:
        return


def _preview_bounds(shapes: list[CadShape], labels: list[CadLabel]) -> tuple[float, float, float, float]:
    points = [point for shape in shapes for point in shape.points]
    points.extend(label.point for label in labels)
    min_x = min(point[0] for point in points)
    min_y = min(point[1] for point in points)
    max_x = max(point[0] for point in points)
    max_y = max(point[1] for point in points)
    if min_x == max_x:
        min_x -= 1
        max_x += 1
    if min_y == max_y:
        min_y -= 1
        max_y += 1
    return min_x, min_y, max_x, max_y


def _make_cad_placeholder_preview(
    path: Path,
    destination: Path,
    label: str,
    notes: list[str] | None = None,
) -> Path | None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1300, 900), "#f8fafc")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default()
    draw.rectangle((40, 40, 1260, 860), outline="#cbd5e1", width=2)
    draw.text((76, 84), "CAD file", fill="#0f172a", font=font)
    draw.text((76, 124), label[:140], fill="#334155", font=font)
    y = 184
    for line in notes or _cad_placeholder_notes(path):
        draw.text((76, y), line[:150], fill="#475569", font=font)
        y += 28
        if y > 780:
            break
    image.save(destination, "JPEG", quality=88, optimize=True)
    return destination


def _cad_placeholder_notes(path: Path) -> list[str]:
    if path.suffix.lower() == ".dwg":
        notes = [
            "DWG preview needs a local DWG to DXF converter.",
            "Supported converters: ODA File Converter, dwgread, or dwg2dxf.",
        ]
        strings = _extract_binary_strings(path, limit=12)
        if strings:
            notes.append("Embedded strings:")
            notes.extend(strings)
        return notes
    return ["No renderable CAD preview was generated."]


def convert_dwg_to_dxf(path: Path, work_dir: Path | None = None) -> Path | None:
    if path.suffix.lower() != ".dwg":
        return path if path.suffix.lower() == ".dxf" else None

    cache_dir = (work_dir or path.parent / ".cad-cache") / "cad-converted"
    cache_dir.mkdir(parents=True, exist_ok=True)
    destination = cache_dir / f"{path.stem}.dxf"
    if destination.exists() and destination.stat().st_mtime >= path.stat().st_mtime:
        return destination

    for converter in (_convert_with_oda, _convert_with_dwgread, _convert_with_dwg2dxf):
        try:
            converted = converter(path, destination, cache_dir)
        except Exception:
            converted = None
        if converted and converted.exists():
            return converted
    return None


def _convert_with_oda(path: Path, destination: Path, cache_dir: Path) -> Path | None:
    executable = shutil.which("ODAFileConverter") or shutil.which("odafileconverter")
    if not executable:
        return None
    subprocess.run(
        [
            executable,
            str(path.parent),
            str(cache_dir),
            "ACAD2018",
            "DXF",
            "0",
            "1",
            path.name,
        ],
        check=False,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=120,
    )
    return _find_converted_dxf(path, destination, cache_dir)


def _convert_with_dwgread(path: Path, destination: Path, cache_dir: Path) -> Path | None:
    executable = shutil.which("dwgread")
    if not executable:
        return None
    subprocess.run(
        [executable, "-O", "DXF", "-o", str(destination), str(path)],
        check=False,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=120,
    )
    return destination if destination.exists() else _find_converted_dxf(path, destination, cache_dir)


def _convert_with_dwg2dxf(path: Path, destination: Path, cache_dir: Path) -> Path | None:
    executable = shutil.which("dwg2dxf")
    if not executable:
        return None
    subprocess.run(
        [executable, "-o", str(destination), str(path)],
        check=False,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=120,
    )
    return destination if destination.exists() else _find_converted_dxf(path, destination, cache_dir)


def _find_converted_dxf(path: Path, destination: Path, cache_dir: Path) -> Path | None:
    candidates = list(cache_dir.rglob(f"{path.stem}.dxf")) + list(cache_dir.rglob(f"{path.stem}.DXF"))
    for candidate in candidates:
        if candidate.resolve() == destination.resolve():
            return destination
        if candidate.exists():
            destination.write_bytes(candidate.read_bytes())
            return destination
    return destination if destination.exists() else None


def _iter_dxf_entities(document, limit: int = 12000):
    count = 0
    for layout in document.layouts:
        try:
            entities = list(layout)
        except Exception:
            continue
        for entity in entities:
            yield entity
            count += 1
            if count >= limit:
                return


def _entity_text(entity) -> str | None:
    entity_type = entity.dxftype()
    value = ""
    if entity_type == "MTEXT":
        value = entity.plain_text()
    elif entity_type in {"TEXT", "ATTRIB", "ATTDEF"}:
        value = getattr(entity.dxf, "text", "")
    elif entity_type == "DIMENSION":
        value = getattr(entity.dxf, "text", "")
        if value in {"", "<>"}:
            try:
                value = f"Dimension measurement {entity.get_measurement():.3f}"
            except Exception:
                value = ""
    return _clean_cad_text(value)


def _entity_insert_point(entity) -> tuple[float, float] | None:
    for attr in ("insert", "defpoint", "location"):
        value = getattr(entity.dxf, attr, None)
        if value is not None:
            return _xy(value)
    return None


def _clean_cad_text(value: str | None) -> str | None:
    if not value:
        return None
    cleaned = re.sub(r"\s+", " ", str(value)).strip()
    cleaned = cleaned.replace("\\P", " ").replace("{", "").replace("}", "")
    return cleaned if re.search(r"[A-Za-z0-9]", cleaned) else None


def _format_counter(counter: Counter[str]) -> str:
    if not counter:
        return "none detected"
    return ", ".join(f"{name} {count}" for name, count in counter.most_common(40))


def _unique_limited(values, limit: int = 120) -> list[str]:
    seen: set[str] = set()
    unique: list[str] = []
    for value in values:
        cleaned = _clean_cad_text(value)
        if not cleaned:
            continue
        key = cleaned.lower()
        if key in seen:
            continue
        seen.add(key)
        unique.append(cleaned)
        if len(unique) >= limit:
            break
    return unique


def _extract_binary_strings(path: Path, limit: int) -> list[str]:
    with path.open("rb") as file:
        data = file.read(MAX_DWG_STRING_SCAN_BYTES)
    values: list[str] = []
    values.extend(match.group().decode("latin-1", errors="ignore") for match in re.finditer(rb"[ -~]{4,}", data))
    values.extend(
        match.group().decode("utf-16le", errors="ignore")
        for match in re.finditer(rb"(?:[\x20-\x7e]\x00){4,}", data)
    )
    return _unique_limited((value for value in values if _looks_like_cad_string(value)), limit)


def _looks_like_cad_string(value: str) -> bool:
    cleaned = value.strip()
    if not 4 <= len(cleaned) <= 160:
        return False
    if not re.search(r"[A-Za-z]", cleaned):
        return False
    if cleaned.count("\\") > 6 or cleaned.count("/") > 10:
        return False
    if len(set(cleaned)) <= 2:
        return False
    return True


def _dwg_header(path: Path) -> str | None:
    with path.open("rb") as file:
        header = file.read(8)
    decoded = header.decode("ascii", errors="ignore").strip("\x00")
    return decoded or None


def _xy(point) -> tuple[float, float]:
    return float(point[0]), float(point[1])


def _circle_points(center: tuple[float, float], radius: float, segments: int = 96) -> list[tuple[float, float]]:
    return [
        (
            center[0] + math.cos(2 * math.pi * index / segments) * radius,
            center[1] + math.sin(2 * math.pi * index / segments) * radius,
        )
        for index in range(segments)
    ]


def _arc_points(
    center: tuple[float, float],
    radius: float,
    start_angle: float,
    end_angle: float,
    segments: int = 72,
) -> list[tuple[float, float]]:
    if end_angle < start_angle:
        end_angle += 360
    steps = max(8, min(segments, int(abs(end_angle - start_angle) / 5) + 1))
    return [
        (
            center[0] + math.cos(math.radians(start_angle + (end_angle - start_angle) * index / steps)) * radius,
            center[1] + math.sin(math.radians(start_angle + (end_angle - start_angle) * index / steps)) * radius,
        )
        for index in range(steps + 1)
    ]


def _layer_color(layer: str) -> str:
    palette = ["#0f766e", "#2563eb", "#d97706", "#7c3aed", "#be123c", "#475569", "#15803d", "#9333ea"]
    return palette[sum(ord(char) for char in layer) % len(palette)]


def validate_stored_upload(path: Path) -> None:
    extension = path.suffix.lower()
    if extension in IMAGE_EXTENSIONS:
        with Image.open(path) as image:
            image.verify()
        return
    if extension == ".pdf":
        PdfReader(str(path))
        return
    if extension == ".docx":
        if not zipfile.is_zipfile(path):
            raise ValueError("Invalid DOCX file")
        return
    if extension in TEXT_EXTENSIONS:
        with path.open("rb") as file:
            file.read(4096).decode("utf-8", errors="ignore")
        return
    if extension == ".dxf":
        import ezdxf

        ezdxf.readfile(path)
        return
    if extension == ".dwg":
        if path.stat().st_size <= 0:
            raise ValueError("Empty DWG file")
        return


def copy_upload(source, destination: Path, max_bytes: int | None = None) -> int:
    destination.parent.mkdir(parents=True, exist_ok=True)
    bytes_written = 0
    with destination.open("wb") as file:
        while chunk := source.read(1024 * 1024):
            bytes_written += len(chunk)
            if max_bytes is not None and bytes_written > max_bytes:
                raise ValueError("File exceeds upload size limit")
            file.write(chunk)
    return bytes_written
